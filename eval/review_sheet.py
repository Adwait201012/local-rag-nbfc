"""Build a side-by-side workbook for judging two runs by eye.

    .venv/bin/python eval/review_sheet.py \\
        --pair untuned eval/nbfc_retrieval_test_questions_results_v1.xlsx eval/nbfc_retrieval_test_questions_results_v2.xlsx \\
        --pair tuned   eval/RBI_NBFC_Retrieval_Questions_results_v1.xlsx eval/RBI_NBFC_Retrieval_Questions_results_v2.xlsx \\
        --out eval/review_v1_vs_v2.xlsx

Rows included:
  * every question whose auto flag changed between the two runs, and
  * a random sample of questions whose flag did not change.

The sample matters. The auto checker matches figures and refusal phrases; it cannot
follow logic. It has already scored an answer that opened with the wrong "Yes" as an
improvement. An unchanged flag is not evidence that the answer is unchanged in quality.

Fill in the yellow Verdict column. The Summary sheet counts the verdicts itself.
"""
from __future__ import annotations

import argparse
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

VERDICTS = ["v1 better", "v2 better", "same", "both wrong"]
FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F4E4A")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="D0D7D6")

COLUMNS = [  # header, width
    ("Set", 10), ("Type", 12), ("Question", 38), ("Expected answer", 34),
    ("v1 answer", 60), ("v1 sources", 24), ("v1 flag", 9),
    ("v2 answer", 60), ("v2 sources", 24), ("v2 flag", 9),
    ("Included because", 22), ("Verdict", 13), ("Notes", 30),
]
VERDICT_COL = "L"


def load(path: str) -> pd.DataFrame:
    df = pd.read_excel(path)
    need = {"question", "type", "actual_answer", "auto_flag"}
    if need - set(df.columns):
        sys.exit(f"{path} is missing {sorted(need - set(df.columns))}; is it a run_sheet results file?")
    return df


def pick_rows(name: str, a: pd.DataFrame, b: pd.DataFrame, sample: int, seed: int) -> list[dict]:
    m = a.merge(b, on="question", suffixes=("_1", "_2"))
    changed = m[m["auto_flag_1"] != m["auto_flag_2"]]
    same = m[m["auto_flag_1"] == m["auto_flag_2"]]
    spot = same.sample(min(sample, len(same)), random_state=seed) if sample else same.iloc[0:0]
    rows = []
    for frame, why in ((changed, None), (spot, "spot check: flag same")):
        for _, r in frame.iterrows():
            rows.append({
                "Set": name,
                "Type": r.get("type_1", ""),
                "Question": r["question"],
                "Expected answer": r.get("expected_answer_1", ""),
                "v1 answer": r["actual_answer_1"],
                "v1 sources": r.get("sources_used_1", ""),
                "v1 flag": r["auto_flag_1"],
                "v2 answer": r["actual_answer_2"],
                "v2 sources": r.get("sources_used_2", ""),
                "v2 flag": r["auto_flag_2"],
                "Included because": why or f"flag {r['auto_flag_1']} -> {r['auto_flag_2']}",
            })
    return rows


def clean(v):
    return "" if v is None or (isinstance(v, float) and pd.isna(v)) else v


def build(rows: list[dict], sets: list[str], out: str) -> None:
    wb = Workbook()

    # ---- Review
    ws = wb.active
    ws.title = "Review"
    for i, (head, width) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=head)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = width
    for r, row in enumerate(rows, 2):
        for i, (head, _) in enumerate(COLUMNS, 1):
            c = ws.cell(row=r, column=i, value=clean(row.get(head, "")))
            c.font = Font(name=FONT, size=9)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        ws[f"{VERDICT_COL}{r}"].fill = INPUT_FILL
        ws[f"M{r}"].fill = INPUT_FILL
        ws.row_dimensions[r].height = 150
    ws.freeze_panes = "D2"
    last = max(len(rows) + 1, 2)
    dv = DataValidation(type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Verdict",
                        error="Choose: " + ", ".join(VERDICTS))
    dv.add(f"{VERDICT_COL}2:{VERDICT_COL}{last}")
    ws.add_data_validation(dv)

    # ---- Summary (formulas, so it updates as verdicts are filled in)
    sm = wb.create_sheet("Summary")
    sm["A1"] = "Verdicts so far"
    sm["A1"].font = Font(name=FONT, bold=True, size=12)
    rng = f"Review!${VERDICT_COL}$2:${VERDICT_COL}${last}"
    setrng = f"Review!$A$2:$A${last}"
    sm.cell(row=3, column=1, value="Verdict")
    sm.cell(row=3, column=2, value="All sets")
    for j, s in enumerate(sets, 3):
        sm.cell(row=3, column=j, value=s)
    for i, v in enumerate(VERDICTS, 4):
        sm.cell(row=i, column=1, value=v)
        sm.cell(row=i, column=2, value=f'=COUNTIF({rng},"{v}")')
        for j, s in enumerate(sets, 3):
            sm.cell(row=i, column=j, value=f'=COUNTIFS({rng},"{v}",{setrng},"{s}")')
    tot = 4 + len(VERDICTS)
    sm.cell(row=tot, column=1, value="Reviewed")
    sm.cell(row=tot, column=2, value=f"=COUNTA({rng})")
    sm.cell(row=tot + 1, column=1, value="Rows to review")
    sm.cell(row=tot + 1, column=2, value=f"=ROWS({rng})")
    for row in sm.iter_rows(min_row=3, max_row=tot + 1):
        for c in row:
            c.font = Font(name=FONT, bold=(c.row == 3 or c.column == 1))
    for col in "ABCDE":
        sm.column_dimensions[col].width = 16

    # ---- How to use
    hw = wb.create_sheet("How to use")
    lines = [
        ("How to review", True),
        ("On the Review sheet, read each question with both answers side by side.", False),
        ("Fill in the yellow Verdict column (drop-down) and, optionally, Notes. Nothing else needs editing.", False),
        ("", False),
        ("Verdicts", True),
        ("v1 better  - you would rather a compliance officer read the v1 answer", False),
        ("v2 better  - you would rather they read the v2 answer", False),
        ("same       - no meaningful difference, including both correct", False),
        ("both wrong - neither answer is acceptable", False),
        ("", False),
        ("What to look for", True),
        ("A missing condition or exception (like tiny deposits) is a real difference.", False),
        ("An answer opening with 'yes' or 'no' that the passages do not support is a serious fault.", False),
        ("A figure or deadline not in the cited passage is the most serious fault of all.", False),
        ("Longer is not better by itself; check what the extra text adds.", False),
        ("", False),
        ("Example", True),
        ("Q: Can a depositor withdraw after two months?  v1: 'Not before three months, except tiny deposits...'", False),
        ("v2: 'Yes, provided the deposit is held for three months...'  ->  Verdict: v1 better  "
         "(v2 opens with a wrong yes)", False),
        ("", False),
        ("Rows marked 'spot check' had the same auto flag in both runs. They are there because the auto "
         "checker cannot follow logic; read them as carefully as the others.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = hw.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT, bold=bold, size=11 if bold else 10)
    hw.column_dimensions["A"].width = 120

    wb.save(out)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--pair", nargs=3, action="append", metavar=("NAME", "V1", "V2"), required=True)
    ap.add_argument("--out", default="eval/review.xlsx")
    ap.add_argument("--sample", type=int, default=5, help="unchanged rows to spot-check per set")
    ap.add_argument("--seed", type=int, default=11)
    args = ap.parse_args()

    rows, sets = [], []
    for name, p1, p2 in args.pair:
        picked = pick_rows(name, load(p1), load(p2), args.sample, args.seed)
        n_changed = sum(1 for r in picked if not r["Included because"].startswith("spot"))
        print(f"{name}: {n_changed} changed, {len(picked) - n_changed} spot checks")
        rows += picked
        sets.append(name)
    build(rows, sets, args.out)
    print(f"\nwritten: {args.out}  ({len(rows)} rows to review)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
